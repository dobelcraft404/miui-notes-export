# -*- coding: utf-8 -*-
"""笔记导出：Markdown、Word(.docx) 与 HTML。

三种输出共用 content.py 解析出的块列表，保证内容一致。
"""

from __future__ import annotations

import csv
import datetime as dt
import html
import os
import re
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from . import media
from .content import ImageBlock, NoteDocument, Span, TextBlock
from .bakfile import Note

FOLDER_ILLEGAL = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
MAX_NAME = 80
MAX_TITLE = 120


_TAG_RE = re.compile(r'<[^>]*>')
_IMAGE_MARK_RE = re.compile(r'\u263a\s*[0-9a-f]{40}(?:\s*<0\s*/>)?')
_WS_RE = re.compile(r'\s+')


def plain_text(text: str) -> str:
    """把标题/摘要里的富文本标记与行内图片引用剥掉，只留可读文字。

    摘要字段（记录 `#8`）和标题字段都可能原样带着 `<text indent="1">`、`<b>`
    这类标记，甚至整段就是一张图片的引用 `☺ <40位hash><0/></>`。
    目录页（index.html / index.csv）直接展示这些字段，不剥就会把标记当正文显示。
    """
    text = _IMAGE_MARK_RE.sub(' ', text or '')
    text = _TAG_RE.sub(' ', text)
    text = text.replace('\u200b', '')
    return _WS_RE.sub(' ', text).strip()


def _clip(text: str, limit: int) -> str:
    """超长文本截断并加省略号。无标题笔记的「标题」会回退成摘要，
    而摘要可能直接是整段正文，不截断会让目录页出现上千字的一行。"""
    if len(text) <= limit:
        return text
    return text[:limit].rstrip(' .,;:，。；：') + '…'


def display_name(note) -> str:
    """笔记的目录名与展示标题：优先标题，其次摘要首行，最后退化为“未命名笔记”。

    摘要可能带富文本标记或直接就是一张图片的引用，必须先把标记剥掉，
    否则会生成 `☺ <hash>...` 这类无意义的目录名。
    """
    for cand in (note.title, note.summary):
        cand = plain_text(cand)
        if cand:
            return _clip(cand, MAX_TITLE)
    return '未命名笔记'


def safe_component(name: str, fallback: str) -> str:
    """净化成单层目录/文件名，并限制长度（Windows 路径上限友好）。"""
    name = FOLDER_ILLEGAL.sub('_', (name or '').replace('\n', ' ')).strip(' .')
    name = re.sub(r'\s+', ' ', name)
    if len(name) > MAX_NAME:
        name = name[:MAX_NAME].rstrip(' .')
    return name or fallback


def fmt_time(ms: int) -> str:
    if not ms or ms < 0:
        return ''
    try:
        return dt.datetime.fromtimestamp(ms / 1000.0).strftime('%Y-%m-%d %H:%M:%S')
    except (OverflowError, OSError, ValueError):
        return str(ms)


def _escape(text: str) -> str:
    # 只转义会干扰 markdown 结构的字符，中文与常见标点原样保留
    return re.sub(r'([\\`*_\[\]])', r'\\\1', text).replace('\u200b', '')


def note_markdown(note: Note, doc: NoteDocument, image_links: Dict[str, str],
                  attachments: Sequence[Tuple[str, str]] = ()) -> str:
    lines: List[str] = ['# ' + display_name(note), '']
    meta = []
    if note.folder:
        meta.append('- 所属目录：%s' % note.folder)
    created, modified = fmt_time(note.created), fmt_time(note.modified)
    if created:
        meta.append('- 创建时间：%s' % created)
    if modified and modified != created:
        meta.append('- 修改时间：%s' % modified)
    meta.append('- 笔记 ID：%s' % note.note_id)
    meta.append('- 类型：%s' % note.kind)
    lines.extend(meta)
    lines.append('')

    body: List[str] = []
    for block in doc.blocks:
        if isinstance(block, ImageBlock):
            link = image_links.get(block.ref)
            if link:
                body.append('![%s](%s)' % (_escape(block.ref), link))
            else:
                body.append('> 附件未在备份中找到：%s' % block.ref)
            body.append('')
            continue
        if not isinstance(block, TextBlock):
            continue
        text = ''.join(
            ('**%s**' % _escape(s.text)) if s.bold else
            ('*%s*' % _escape(s.text)) if s.italic else
            _escape(s.text)
            for s in block.spans
        )
        if block.checkbox is not None:
            # 待办项：缩进用两个空格，保持 markdown 列表层级
            body.append(('  ' * min(block.indent, 4)) +
                        ('- [%s] %s' % ('x' if block.checkbox else ' ', text)))
            continue
        if not text.strip():
            body.append('')
            continue
        # 小米笔记的缩进层级用引用符号表达（与 App 内的视觉层级一致）
        body.append(('> ' * min(block.indent, 3)) + text)
    # 折叠连续空行
    compact: List[str] = []
    for line in body:
        if line == '' and compact and compact[-1] == '':
            continue
        compact.append(line)
    lines.extend(compact)
    lines.append('')

    for ref, rel in attachments:
        lines.append('- 附件文件：`%s`（%s）' % (rel, ref))
    if attachments:
        lines.append('')
    return '\n'.join(lines).rstrip() + '\n'


# ---------------------------------------------------------------------------
# Word 导出
# ---------------------------------------------------------------------------

def _add_spans(paragraph, spans: Sequence[Span]) -> None:
    for s in spans:
        run = paragraph.add_run(s.text)
        run.bold = s.bold
        run.italic = s.italic
        run.underline = s.underline


def _add_image(document, path: str, max_width_cm: float = 15.0):
    from docx.shared import Cm
    try:
        from PIL import Image  # 可选依赖，用于等比缩放
    except Exception:
        Image = None
    width = None
    if Image is not None:
        try:
            with Image.open(path) as im:
                w, h = im.size
            if w > 0:
                width = Cm(max_width_cm)
                if h / w > 1.4:
                    width = Cm(min(max_width_cm, max_width_cm * 0.75))
        except Exception:
            width = None
    try:
        document.add_picture(path, width=width)
    except Exception:
        return False
    return True


def note_docx(note: Note, doc: NoteDocument, image_paths: Dict[str, str],
              out_path: str, attachments: Sequence[Tuple[str, str]] = ()) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    style = document.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    # 中文字体需要单独设置 eastAsia，否则 Word 会回退到宋体
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
    if rfonts is None:
        from docx.oxml.ns import qn
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    from docx.oxml.ns import qn
    rfonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    document.add_heading(display_name(note), level=1)
    info = []
    if note.folder:
        info.append('所属目录：%s' % note.folder)
    if note.created:
        info.append('创建时间：%s' % fmt_time(note.created))
    if note.modified and note.modified != note.created:
        info.append('修改时间：%s' % fmt_time(note.modified))
    info.append('笔记 ID：%s' % note.note_id)
    info.append('类型：%s' % note.kind)
    p = document.add_paragraph('　|　'.join(info))
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    for block in doc.blocks:
        if isinstance(block, ImageBlock):
            path = image_paths.get(block.ref)
            if path and os.path.exists(path):
                _add_image(document, path)
                cap = document.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(os.path.basename(path))
                run.font.size = Pt(8)
                run.italic = True
            else:
                document.add_paragraph('［附件缺失：%s］' % block.ref)
            continue
        if not isinstance(block, TextBlock):
            continue
        if block.checkbox is not None:
            para = document.add_paragraph()
            para.paragraph_format.left_indent = None
            mark = '☑ ' if block.checkbox else '☐ '
            if block.indent:
                from docx.shared import Cm
                para.paragraph_format.left_indent = Cm(0.6 * min(block.indent, 8))
            _add_spans(para, [Span(mark)] + list(block.spans))
            continue
        para = document.add_paragraph()
        if block.indent:
            from docx.shared import Cm
            para.paragraph_format.left_indent = Cm(0.6 * min(block.indent, 8))
        text = block.text
        if block.indent and text.strip() and all(not s.bold for s in block.spans) \
                and len(text) < 40 and not para.text:
            # 短且无格式的缩进行视为小标题
            heading = document.add_heading(level=min(2 + block.indent, 4))
            heading.text = ''
            _add_spans(heading, block.spans)
            continue
        _add_spans(para, block.spans)

    if attachments:
        document.add_heading('附件清单', level=2)
        for ref, rel in attachments:
            document.add_paragraph('• %s' % rel)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    document.save(out_path)


# ---------------------------------------------------------------------------
# HTML 导出
# ---------------------------------------------------------------------------

CSS = """
:root { color-scheme: light dark; }
* { box-sizing: border-box; }
body {
  margin: 0; padding: 2rem 1rem 4rem;
  font-family: "Microsoft YaHei", "PingFang SC", "Noto Sans CJK SC", system-ui, sans-serif;
  font-size: 15px; line-height: 1.75; color: #24292f; background: #f6f8fa;
}
.page { max-width: 46rem; margin: 0 auto; background: #fff; border-radius: 10px;
        padding: 2rem 2.25rem 2.5rem; box-shadow: 0 1px 4px rgba(0,0,0,.08); }
h1 { font-size: 1.5rem; margin: 0 0 .6rem; line-height: 1.4; }
.meta { color: #6a737d; font-size: .82rem; margin: 0 0 1.6rem;
        padding-bottom: .9rem; border-bottom: 1px solid #eaecef; }
.meta span + span::before { content: " · "; }
p { margin: .55rem 0; }
p.indent-1 { margin-left: 1.4rem; }
p.indent-2 { margin-left: 2.8rem; }
p.indent-3 { margin-left: 4.2rem; }
.todo { list-style: none; padding-left: 0; margin: .4rem 0; }
.todo li { margin: .25rem 0; }
.todo input { margin-right: .5rem; vertical-align: -1px; }
figure { margin: 1rem 0; }
figure img { max-width: 100%; height: auto; border-radius: 6px; border: 1px solid #eaecef; }
figcaption { color: #8b949e; font-size: .75rem; margin-top: .35rem; word-break: break-all; }
.attachments { margin-top: 2rem; padding-top: .9rem; border-top: 1px dashed #d8dee4;
               color: #57606a; font-size: .84rem; }
.attachments h2 { font-size: .95rem; margin: 0 0 .5rem; color: #24292f; }
.attachments ul { margin: 0; padding-left: 1.2rem; }
.attachments code { word-break: break-all; }
a { color: #0969da; }
.index-list { list-style: none; padding: 0; margin: 0; }
.index-list li { padding: .6rem 0; border-bottom: 1px solid #eaecef; }
.index-list li:last-child { border-bottom: 0; }
.index-list .title { font-weight: 600; }
.index-list .sub { color: #6a737d; font-size: .8rem; }
.index-list .links a + a::before { content: " | "; color: #d0d7de; }
.group h2 { font-size: 1.05rem; margin: 1.8rem 0 .4rem; padding-bottom: .3rem;
            border-bottom: 2px solid #eaecef; }
#q { width: 100%; padding: .55rem .8rem; font-size: .95rem; margin: .2rem 0 1.2rem;
     border: 1px solid #d0d7de; border-radius: 6px; font-family: inherit; }
@media (prefers-color-scheme: dark) {
  body { background: #0d1117; color: #c9d1d9; }
  .page { background: #161b22; box-shadow: none; }
  .meta, .index-list .sub, figcaption { color: #8b949e; }
  .meta, .index-list li, .group h2 { border-color: #30363d; }
  .attachments { border-color: #30363d; color: #8b949e; }
  .attachments h2 { color: #c9d1d9; }
  #q { background: #0d1117; color: #c9d1d9; border-color: #30363d; }
  a { color: #58a6ff; }
  figure img { border-color: #30363d; }
}
"""


def _esc(text: str) -> str:
    return html.escape(text or '', quote=True)


def _span_html(spans: Sequence[Span]) -> str:
    out: List[str] = []
    for s in spans:
        piece = _esc(s.text)
        if s.bold:
            piece = '<strong>%s</strong>' % piece
        if s.italic:
            piece = '<em>%s</em>' % piece
        if s.underline:
            piece = '<u>%s</u>' % piece
        out.append(piece)
    return ''.join(out)


def _indent_class(indent: int) -> str:
    return (' indent-%d' % min(indent, 3)) if indent else ''


def _html_document(title: str, body: str) -> str:
    return ('<!DOCTYPE html>\n<html lang="zh-CN">\n<head>\n'
            '<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
            '<title>%s</title>\n<style>%s</style>\n</head>\n<body>\n%s\n</body>\n</html>\n'
            % (_esc(title), CSS, body))


def note_html(note: Note, doc: NoteDocument, image_links: Dict[str, str],
              attachments: Sequence[Tuple[str, str]] = ()) -> str:
    """单篇笔记的完整 HTML 页面。图片按相对路径引用 note.assets/。"""
    meta = []
    if note.folder:
        meta.append('<span>目录：%s</span>' % _esc(note.folder))
    created, modified = fmt_time(note.created), fmt_time(note.modified)
    if created:
        meta.append('<span>创建：%s</span>' % _esc(created))
    if modified and modified != created:
        meta.append('<span>修改：%s</span>' % _esc(modified))
    meta.append('<span>笔记 ID：%s</span>' % _esc(note.note_id))
    meta.append('<span>类型：%s</span>' % _esc(note.kind))

    body: List[str] = ['<div class="page">',
                       '<h1>%s</h1>' % _esc(display_name(note)),
                       '<p class="meta">%s</p>' % ''.join(meta)]

    open_todo: List[str] = []

    def close_todo() -> None:
        if open_todo:
            body.append('<ul class="todo">%s</ul>' % ''.join(open_todo))
            open_todo.clear()

    for block in doc.blocks:
        if isinstance(block, ImageBlock):
            close_todo()
            link = image_links.get(block.ref)
            if link:
                caption = block.caption or block.ref
                body.append('<figure><img src="%s" alt="%s" loading="lazy">'
                            '<figcaption>%s</figcaption></figure>'
                            % (_esc(link), _esc(caption), _esc(caption)))
            else:
                body.append('<p>［附件未在备份中找到：%s］</p>' % _esc(block.ref))
            continue
        if not isinstance(block, TextBlock):
            continue
        if block.checkbox is not None:
            checked = ' checked' if block.checkbox else ''
            inner = _span_html(block.spans)
            style = '' if not block.indent else ' style="margin-left:%.1fem"' % (1.4 * min(block.indent, 4))
            open_todo.append('<li%s><label><input type="checkbox" disabled%s>%s</label></li>'
                             % (style, checked, inner))
            continue
        close_todo()
        if not block.text.strip():
            body.append('<p>&nbsp;</p>')
            continue
        body.append('<p class="para%s">%s</p>'
                    % (_indent_class(block.indent), _span_html(block.spans) or '&nbsp;'))
    close_todo()

    if attachments:
        items = ''.join('<li><code>%s</code>（%s）</li>' % (_esc(rel), _esc(ref))
                        for ref, rel in attachments)
        body.append('<div class="attachments"><h2>附件清单</h2><ul>%s</ul></div>' % items)
    body.append('</div>')
    return _html_document(display_name(note), '\n'.join(body))


def write_index_html(path: str, rows: Sequence[dict],
                     formats: Sequence[str] = ('md', 'docx', 'html')) -> None:
    """生成可直接在浏览器里翻阅的目录页（含搜索框，纯静态、可离线打开）。"""
    groups: Dict[str, List[dict]] = {}
    for row in rows:
        groups.setdefault(row.get('folder') or '未分类', []).append(row)

    parts: List[str] = ['<div class="page">',
                        '<h1>小米笔记导出</h1>',
                        '<p class="meta"><span>共 %d 篇笔记</span><span>%d 个目录</span></p>'
                        % (len(rows), len(groups)),
                        '<input id="q" type="search" placeholder="输入关键词过滤标题 / 目录…" '
                        'autocomplete="off">']

    def note_links(row: dict) -> str:
        links = []
        base = row['path'].replace('\\', '/') if row.get('path') else ''
        if base:
            for fmt in formats:
                if row.get(fmt):
                    links.append('<a href="%s/%s">%s</a>'
                                 % (_esc(base), _esc(row[fmt]), fmt.upper()))
        return ''.join(links)

    for folder in sorted(groups):
        parts.append('<div class="group"><h2>%s</h2><ul class="index-list">' % _esc(folder))
        for row in sorted(groups[folder], key=lambda r: r.get('modified') or '', reverse=True):
            bits = []
            if row.get('created'):
                bits.append('创建 %s' % row['created'])
            if row.get('kind') and row['kind'] != 'common':
                bits.append(row['kind'])
            bits.append('%s 图 / %s 附件' % (row.get('images', 0), row.get('attachments', 0)))
            parts.append(
                '<li data-text="%s"><span class="title">%s</span>'
                ' <span class="links">%s</span><br>'
                '<span class="sub">%s</span></li>'
                % (_esc('%s %s' % (folder, row.get('title', ''))),
                   _esc(row.get('title') or '未命名笔记'), note_links(row),
                   _esc(' · '.join(bits))))
        parts.append('</ul></div>')
    parts.append('</div>')

    script = """
<script>
(function () {
  var q = document.getElementById('q');
  var items = [].slice.call(document.querySelectorAll('.index-list li'));
  var groups = [].slice.call(document.querySelectorAll('.group'));
  q.addEventListener('input', function () {
    var needle = q.value.trim().toLowerCase();
    items.forEach(function (li) {
      li.style.display = !needle ||
        (li.getAttribute('data-text') || '').toLowerCase().indexOf(needle) >= 0 ? '' : 'none';
    });
    groups.forEach(function (g) {
      var any = [].slice.call(g.querySelectorAll('.index-list li'))
        .some(function (li) { return li.style.display !== 'none'; });
      g.style.display = any ? '' : 'none';
    });
  });
})();
</script>"""

    body = '\n'.join(parts) + script
    os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write(_html_document('小米笔记导出', body))


# ---------------------------------------------------------------------------
# 批量导出
# ---------------------------------------------------------------------------

def write_index_csv(path: str, rows: Iterable[dict]) -> None:
    fields = ['note_id', 'folder', 'title', 'kind', 'created', 'modified',
              'blocks', 'images', 'attachments', 'md', 'docx', 'html']
    os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
    with open(path, 'w', encoding='utf-8-sig', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction='ignore')
        w.writeheader()
        for r in rows:
            w.writerow(r)
